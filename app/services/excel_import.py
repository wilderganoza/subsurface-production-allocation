"""Lectura de los Excel de producción, petrofísica e intervenciones.

Porta client/src/utils/fileParser.js. Moverlo al servidor evita que el navegador
tenga que cargar la librería xlsx y permite validar los datos antes de guardarlos.

Los tres formatos son los de examples/:
  - Producción:    Date | Oil Rate (BOPD)
  - Petrofísica:   Sand | kh (md·ft)
  - Intervenciones: Sand | <fecha> | <fecha> | ...  con "X" donde la arena está abierta
"""

# Importamos las librerias necesarias
import io  # Para leer el archivo desde memoria
from datetime import date, datetime  # Para normalizar fechas
from typing import Any, Dict, List, Tuple  # Tipos

from openpyxl import load_workbook  # Lectura de .xlsx

from app.core.logging import get_logger  # Para registrar filas descartadas

# Creamos el logger de este módulo
logger = get_logger(__name__)


# Definimos el error que lanzamos cuando el archivo no tiene el formato esperado
class ImportError_(Exception):
    """El archivo no se pudo interpretar con el formato esperado."""


# Normalizamos un valor de celda a fecha
def _to_date(value: Any) -> date:
    # Un datetime se reduce a su fecha
    if isinstance(value, datetime):
        return value.date()

    # Una fecha se devuelve tal cual
    if isinstance(value, date):
        return value

    # El texto se parsea en ISO, que es como lo escriben los Excel de ejemplo
    texto = str(value).strip()[:10]

    # Dejamos que el error suba con contexto si el formato no es reconocible
    try:
        # Intentamos el formato ISO
        return datetime.fromisoformat(texto).date()
    # Un valor no interpretable es un error del archivo, no del programa
    except ValueError as exc:
        raise ImportError_(f"No se pudo interpretar «{value}» como fecha") from exc


# Normalizamos un valor de celda a número
def _to_float(value: Any) -> float:
    # Una celda vacía cuenta como cero
    if value is None or str(value).strip() == "":
        return 0.0

    # Aceptamos coma decimal, que es lo habitual en los Excel en español
    texto = str(value).strip().replace(",", ".")

    # Dejamos que el error suba con contexto
    try:
        # Convertimos a número
        return float(texto)
    # Un valor no numérico es un error del archivo
    except ValueError as exc:
        raise ImportError_(f"No se pudo interpretar «{value}» como número") from exc


# Abrimos el libro y devolvemos las filas de la primera hoja
def _read_rows(content: bytes) -> List[Tuple]:
    # Abrimos desde memoria en modo solo lectura
    try:
        # data_only para tomar los valores calculados, no las fórmulas
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    # Un archivo corrupto o que no es xlsx
    except Exception as exc:
        raise ImportError_(f"El archivo no es un Excel válido: {exc}") from exc

    # Tomamos la primera hoja, que es donde viven los datos en los tres formatos
    sheet = workbook[workbook.sheetnames[0]]

    # Materializamos las filas antes de cerrar el libro
    rows = list(sheet.iter_rows(values_only=True))

    # Cerramos para liberar el archivo
    workbook.close()

    # Sin cabecera y al menos una fila no hay nada que importar
    if len(rows) < 2:
        raise ImportError_("El archivo no tiene filas de datos")

    # Devolvemos todas las filas, cabecera incluida
    return rows


# Leemos el histórico de producción
def parse_production(content: bytes) -> List[Dict[str, Any]]:
    # Recuperamos las filas
    rows = _read_rows(content)

    # Acumulamos los registros válidos
    records = []

    # Recorremos saltando la cabecera
    for i, row in enumerate(rows[1:], start=2):
        # Descartamos las filas sin fecha o sin valor
        if not row or row[0] is None or len(row) < 2 or row[1] is None:
            continue

        # Convertimos ambos campos, indicando la fila si algo falla
        try:
            # La fecha va en la primera columna y la producción en la segunda
            records.append({
                "date": _to_date(row[0]),
                "total_production": _to_float(row[1]),
            })
        # Enriquecemos el error con el número de fila, que es lo que el usuario ve
        except ImportError_ as exc:
            raise ImportError_(f"Fila {i}: {exc}") from exc

    # Sin registros el archivo no aporta nada
    if not records:
        raise ImportError_("No se encontró ningún dato de producción")

    # Ordenamos por fecha: el motor asume orden cronológico
    records.sort(key=lambda r: r["date"])

    # Devolvemos el histórico
    return records


# Leemos las propiedades petrofísicas
def parse_petrophysics(content: bytes) -> List[Dict[str, Any]]:
    # Recuperamos las filas
    rows = _read_rows(content)

    # Acumulamos las arenas válidas
    sands = []

    # Recorremos saltando la cabecera
    for i, row in enumerate(rows[1:], start=2):
        # Descartamos las filas sin nombre de arena
        if not row or row[0] is None or str(row[0]).strip() == "":
            continue

        # Convertimos el k·h, indicando la fila si falla
        try:
            # El nombre va en la primera columna y el k·h en la segunda
            kh = _to_float(row[1] if len(row) > 1 else 0)
        # Enriquecemos el error con el número de fila
        except ImportError_ as exc:
            raise ImportError_(f"Fila {i}: {exc}") from exc

        # Un k·h negativo no tiene sentido físico
        if kh < 0:
            raise ImportError_(f"Fila {i}: el k·h no puede ser negativo")

        # Guardamos la arena
        sands.append({"sand_name": str(row[0]).strip(), "kh": kh})

    # Sin arenas no se puede repartir nada
    if not sands:
        raise ImportError_("No se encontró ninguna arena")

    # Devolvemos las arenas
    return sands


# Leemos la matriz de intervenciones
def parse_interventions(content: bytes) -> Dict[str, Any]:
    # Recuperamos las filas
    rows = _read_rows(content)

    # La cabecera lleva las fechas a partir de la segunda columna
    header = rows[0]

    # Convertimos las fechas de la cabecera
    dates = []

    # Recorremos las columnas de la cabecera
    for j, cell in enumerate(header[1:], start=2):
        # Paramos en la primera columna vacía
        if cell is None or str(cell).strip() == "":
            break

        # Convertimos, indicando la columna si falla
        try:
            # Cada columna es una fecha de intervención
            dates.append(_to_date(cell))
        # Enriquecemos el error con la columna
        except ImportError_ as exc:
            raise ImportError_(f"Columna {j} de la cabecera: {exc}") from exc

    # Sin fechas no hay periodos que construir
    if not dates:
        raise ImportError_("La cabecera no contiene fechas de intervención")

    # Acumulamos las arenas y la matriz
    sand_names = []
    matrix = []

    # Recorremos saltando la cabecera
    for row in rows[1:]:
        # Descartamos las filas sin nombre de arena
        if not row or row[0] is None or str(row[0]).strip() == "":
            continue

        # Guardamos el nombre
        sand_names.append(str(row[0]).strip())

        # Cualquier contenido en la celda marca la arena como abierta
        matrix.append([
            bool(row[j] and str(row[j]).strip()) if j < len(row) else False
            for j in range(1, len(dates) + 1)
        ])

    # Sin arenas la matriz no sirve
    if not sand_names:
        raise ImportError_("No se encontró ninguna arena en la matriz")

    # Devolvemos la matriz en la forma que espera el motor de asignación
    return {"sandNames": sand_names, "interventionDates": dates, "matrix": matrix}
